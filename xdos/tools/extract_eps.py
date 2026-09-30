#!/usr/bin/env python3
"""
extract_eps.py — Extract EPS (NL EPS Communication Protocol) commands
and their responses from the protocol workbook.

Input  : the .xlsx workbook (one sheet per command)
Output : a dict of records consumed by build_mib.py

EPS frame on the wire:  [dest][src][msgID][data...][crc8]
Each command sheet has a Request block (-> TC) and a Response block
(-> TM), laid out side by side:

    Request                          |  Response
    Index | Parameter | Value | Cmt  |  Index | Parameter | Value | Cmt

The data fields stored as MIB parameters are the payload bytes only
(index >= 3, excluding the trailing CRC); offsets are re-based so
that payload byte 0 has offset 0. Framing bytes (dest/src/msgID/crc)
are handled by the codec, not stored as parameters.
"""

from __future__ import annotations
import re
from dataclasses import dataclass, field
from typing import Optional

try:
    import openpyxl
except ImportError:
    openpyxl = None

# sheets that are not command definitions
NON_COMMAND_SHEETS = {
    'Command_Status', 'Notes', 'Msg Structure', 'Change_Log',
}

# representative RS485 addresses (operator can re-address per physical board)
BOARD_ADDR = {'PDU': 1, 'BMS': 9, 'MPPT': 12}

FRAMING_PARAMS = re.compile(
    r'(device\s+address|message\s+id|crc)', re.IGNORECASE)


@dataclass
class Param:
    seq: int
    offset: int          # bytes, re-based to payload start
    length: int          # bytes
    name: str
    data_type: str
    description: str


@dataclass
class Message:
    name: str
    direction: str       # 'TC' | 'TM'
    msg_id: int
    board: Optional[str] # 'PDU' | 'BMS' | 'MPPT' | None (all boards)
    length_bytes: Optional[int]
    description: str = ''
    doc_reference: str = ''
    params: list = field(default_factory=list)


def _board_of(sheet_name: str) -> Optional[str]:
    s = sheet_name.upper()
    if s.startswith('PDU'):
        return 'PDU'
    if s.startswith('BMS'):
        return 'BMS'
    if s.startswith('MPPT'):
        return 'MPPT'
    return None          # 'All boards' commands (FW version, bootloader, ...)


def _parse_index(raw) -> Optional[tuple]:
    """Return (offset_bytes, length_bytes) from a doc index cell.
    Accepts '7' or '3-6'. Offsets are re-based: payload byte 0 = doc index 3."""
    if raw is None:
        return None
    s = str(raw).strip()
    m = re.fullmatch(r'(\d+)\s*-\s*(\d+)', s)
    if m:
        lo, hi = int(m.group(1)), int(m.group(2))
        return (lo - 3, hi - lo + 1)
    if s.isdigit():
        return (int(s) - 3, 1)
    return None


def _infer_type(length: int, comment: str, name: str) -> str:
    c = (comment or '').lower()
    n = (name or '').lower()
    if 'float' in c or 'float' in n:
        return 'FLOAT'
    if re.search(r'\b\d+\s*:\s*\w', comment or ''):   # "0: NS 1: NL ..."
        return 'ENUM'
    if length >= 4:
        return 'UINT'        # 4-byte non-float integer
    if length == 2:
        return 'UINT'
    return 'UINT'            # single byte


def _block_rows(ws, idx_col: int, max_col_excl: int):
    """Yield (index_cell, param_cell, value_cell, comment_cell) for a block,
    starting after the 'Index' header row."""
    header_row = None
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cell = row[idx_col] if idx_col < len(row) else None
        if cell is not None and str(cell).strip().lower() == 'index':
            header_row = r
            break
    if header_row is None:
        return
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        idx = row[idx_col] if idx_col < len(row) else None
        prm = row[idx_col + 1] if idx_col + 1 < len(row) else None
        val = row[idx_col + 2] if idx_col + 2 < len(row) else None
        cmt = row[idx_col + 3] if idx_col + 3 < min(len(row), max_col_excl) else None
        if idx is None and prm is None:
            continue
        yield idx, prm, val, cmt


def _find_response_col(ws) -> Optional[int]:
    """Locate the column where the Response block's 'Index' header sits."""
    for row in ws.iter_rows(values_only=True):
        for c, cell in enumerate(row):
            if cell is not None and str(cell).strip() == 'Response':
                # the 'Index' header is at or just after this column
                return c
    return None


def _extract_block(ws, idx_col, max_col_excl, direction, sheet_name):
    msg = Message(name=sheet_name, direction=direction, msg_id=-1,
                  board=_board_of(sheet_name), length_bytes=None,
                  doc_reference=sheet_name)
    seq = 0
    last_payload_end = 0
    for idx, prm, val, cmt in _block_rows(ws, idx_col, max_col_excl):
        pname = str(prm).strip() if prm is not None else ''
        # capture message id
        if pname.lower() == 'message id' and val is not None:
            try:
                msg.msg_id = int(val)
            except (TypeError, ValueError):
                pass
            continue
        # skip framing rows
        if FRAMING_PARAMS.search(pname):
            continue
        if not pname:
            continue
        span = _parse_index(idx)
        if span is None:
            continue
        offset, length = span
        if offset < 0:
            continue
        seq += 1
        dtype = _infer_type(length, str(cmt or ''), pname)
        msg.params.append(Param(seq, offset, length, pname, dtype,
                                str(cmt).strip() if cmt is not None else ''))
        last_payload_end = max(last_payload_end, offset + length)
    if msg.params:
        msg.length_bytes = last_payload_end
    return msg if msg.msg_id >= 0 else None


def extract(xlsx_path: str) -> dict:
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    messages = []
    skipped = []
    for sheet_name in wb.sheetnames:
        if sheet_name in NON_COMMAND_SHEETS:
            continue
        ws = wb[sheet_name]
        resp_col = _find_response_col(ws)
        # Request block is column 0; Response block starts at resp_col
        req = _extract_block(ws, 0, resp_col or ws.max_column, 'TC', sheet_name)
        tm = None
        if resp_col is not None:
            tm = _extract_block(ws, resp_col, ws.max_column + 1, 'TM', sheet_name)
        if req:
            messages.append(req)
        else:
            skipped.append(sheet_name + ' (request)')
        if tm:
            messages.append(tm)
    return {'subsystem': 'EPS', 'protocol': 'eps_rs485',
            'offset_unit': 'bytes', 'board_addr': BOARD_ADDR,
            'messages': messages, 'enums': [], 'skipped': skipped}


if __name__ == '__main__':
    import sys
    r = extract(sys.argv[1] if len(sys.argv) > 1 else '/tmp/eps.xlsx')
    tc = sum(1 for m in r['messages'] if m.direction == 'TC')
    tm = sum(1 for m in r['messages'] if m.direction == 'TM')
    np = sum(len(m.params) for m in r['messages'])
    print(f"EPS: {tc} TC, {tm} TM, {np} params, {len(r['skipped'])} skipped")
    for m in r['messages'][:6]:
        print(f"  [{m.direction}] {m.name} id={m.msg_id} board={m.board} "
              f"len={m.length_bytes} params={len(m.params)}")
    for s in r['skipped']:
        print('   skipped:', s)
